from io import BytesIO

import pandas as pd
from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from vnstock3 import Vnstock
except ImportError:
    from vnstock import Vnstock

from config import DEVELOPER_EMAIL, DEVELOPER_GITHUB, DEVELOPER_WEB
from constants import RISK_WARNING, SOURCE
from utils.dataframe_utils import prepare_ohlcv_dataframe
from utils.validators import MAX_EXPORT_ROWS, validate_date_range, validate_symbol


NUMBER_FORMAT_PRICE = '#,##0'
NUMBER_FORMAT_PERCENT = '0.00'
NUMBER_FORMAT_VOLUME = '#,##0'


def get_stock_api(symbol_upper):
    return Vnstock().stock(symbol=symbol_upper, source=SOURCE)


def create_excel_report(symbol, start_date, end_date):
    try:
        symbol_upper, symbol_error = validate_symbol(symbol)
        if symbol_error:
            return None, None, {"error": symbol_error}, 400

        start_obj, end_obj, date_error = validate_date_range(start_date, end_date)
        if date_error:
            return None, None, {"error": date_error}, 400

        try:
            stock_api = get_stock_api(symbol_upper)
            raw_df = stock_api.quote.history(start=start_date, end=end_date, interval="1D")
        except Exception:
            current_app.logger.exception("Cannot fetch export history for %s", symbol_upper)
            return None, None, {"error": "Không lấy được dữ liệu chứng khoán từ nguồn bên ngoài. Vui lòng thử lại sau."}, 500

        df = prepare_ohlcv_dataframe(raw_df, multiply_price_by_1000=True)

        if df is None or df.empty:
            return None, None, {"error": f"Không có dữ liệu lịch sử để xuất file cho mã {symbol_upper}."}, 404

        df = df[(df["time"] >= pd.to_datetime(start_obj)) & (df["time"] <= pd.to_datetime(end_obj))].reset_index(drop=True)

        if df.empty:
            return None, None, {"error": f"Không có dữ liệu trong khoảng ngày đã chọn cho mã {symbol_upper}."}, 404

        if len(df) > MAX_EXPORT_ROWS:
            return None, None, {"error": f"Dữ liệu export vượt quá {MAX_EXPORT_ROWS} dòng. Vui lòng chọn khoảng ngày ngắn hơn."}, 400

        df["Change"] = df["close"].diff().fillna(0)
        df["PercentChange"] = (df["close"].pct_change() * 100).replace([float("inf"), float("-inf")], 0).fillna(0)

        wb = Workbook()
        ws = wb.active
        ws.title = symbol_upper[:31]

        normal_font = Font(name="Arial", size=10)
        title_font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        subtitle_font = Font(name="Arial", size=10, bold=True)
        header_font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        disclaimer_font = Font(name="Arial", size=9, bold=True, color="CC0000")
        title_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A1:H1")
        ws["A1"] = f"VSDAT - Báo cáo dữ liệu lịch sử {symbol_upper}"
        ws["A1"].font = title_font
        ws["A1"].fill = title_fill
        ws["A1"].alignment = center

        ws.merge_cells("A2:H2")
        ws["A2"] = f"Khoảng thời gian: {start_date} đến {end_date}"
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = left

        ws.merge_cells("A3:H3")
        ws["A3"] = RISK_WARNING
        ws["A3"].font = disclaimer_font
        ws["A3"].alignment = left

        if DEVELOPER_EMAIL:
            ws["A4"] = "Email"
            ws["B4"] = DEVELOPER_EMAIL
        ws["C4"] = "GitHub"
        ws["D4"] = "Tại đây"
        ws["D4"].hyperlink = DEVELOPER_GITHUB
        ws["D4"].style = "Hyperlink"
        ws["E4"] = "Web"
        ws["F4"] = "Tại đây"
        ws["F4"].hyperlink = DEVELOPER_WEB
        ws["F4"].style = "Hyperlink"

        headers = [
            "NGÀY", "GIÁ MỞ CỬA", "GIÁ CAO NHẤT", "GIÁ THẤP NHẤT",
            "GIÁ ĐÓNG CỬA", "THAY ĐỔI GIÁ", "% THAY ĐỔI", "KHỐI LƯỢNG",
        ]
        header_row = 6
        data_start_row = header_row + 1

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

        for row_num, (_, row) in enumerate(df.iterrows(), data_start_row):
            values = [
                row["time"].strftime("%d/%m/%Y"),
                float(row["open"]),
                float(row["high"]),
                float(row["low"]),
                float(row["close"]),
                float(row["Change"]),
                float(row["PercentChange"]),
                float(row["volume"]),
            ]

            for col_num, value in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.font = normal_font
                cell.alignment = center

                if col_num in [2, 3, 4, 5, 6]:
                    cell.number_format = NUMBER_FORMAT_PRICE
                elif col_num == 7:
                    cell.number_format = NUMBER_FORMAT_PERCENT
                elif col_num == 8:
                    cell.number_format = NUMBER_FORMAT_VOLUME

        for col_num in range(1, 9):
            column_letter = get_column_letter(col_num)
            max_length = 12
            for cell in ws[column_letter]:
                if cell.value is not None:
                    max_length = max(max_length, min(len(str(cell.value)) + 2, 50))
            ws.column_dimensions[column_letter].width = max_length

        ws.freeze_panes = f"A{data_start_row}"
        ws.auto_filter.ref = f"A{header_row}:H{header_row}"

        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        filename = f"{symbol_upper}_{start_date}_{end_date}_VSDAT.xlsx"
        return excel_file, filename, None, 200

    except Exception:
        current_app.logger.exception("Excel export failed")
        return None, None, {"error": "Lỗi hệ thống khi xuất Excel. Vui lòng thử lại sau."}, 500
