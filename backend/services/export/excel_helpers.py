from __future__ import annotations

from copy import copy
from typing import Any, Optional

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries


def _clean_sheet_values(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
    if src.hyperlink:
        dst._hyperlink = copy(src.hyperlink)
    if src.comment:
        dst.comment = copy(src.comment)


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        _copy_cell_style(ws.cell(source_row, col), ws.cell(target_row, col))
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _ensure_template_rows(ws, source_row: int, start_row: int, row_count: int, max_col: int) -> None:
    for row in range(start_row, start_row + row_count):
        if row > ws.max_row:
            ws.append([None] * max_col)
        _copy_row_style(ws, source_row, row, max_col)


def _unmerge_ranges_from_row(ws, min_row: int) -> None:
    """
    Template có một số merge cell cố định cho mẫu 5 tháng.
    Khi xuất 1-10 năm, bảng monthly summary dài hơn và có thể đi xuống đúng
    vùng đang merge, làm openpyxl báo: MergedCell object attribute value is read-only.
    Vì vậy chỉ giữ merge tiêu đề phía trên, còn vùng dữ liệu động từ min_row trở xuống
    thì unmerge để có thể ghi dữ liệu dài hạn an toàn.
    """
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= min_row:
            ws.unmerge_cells(str(merged_range))


def _set_value(ws, row: int, col: int, value: Any) -> None:
    """Ghi giá trị an toàn, tự bỏ merge nếu ô đích là MergedCell."""
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        for merged_range in list(ws.merged_cells.ranges):
            if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
                ws.unmerge_cells(str(merged_range))
                break
    ws.cell(row, col).value = value


def _set_series_formula(series, value_formula: Optional[str] = None, category_formula: Optional[str] = None) -> None:
    if value_formula and getattr(series, "val", None) and getattr(series.val, "numRef", None):
        series.val.numRef.f = value_formula

    if category_formula and getattr(series, "cat", None):
        if getattr(series.cat, "strRef", None):
            series.cat.strRef.f = category_formula
        elif getattr(series.cat, "numRef", None):
            series.cat.numRef.f = category_formula


def _set_chart_series_title(series, title: str) -> None:
    try:
        series.tx.v = title
    except Exception:
        pass


def _set_workbook_calc_mode(wb) -> None:
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def _center_range(ws, cell_range: str) -> None:
    """
    Chỉ căn giữa nội dung trong vùng bảng, không thay đổi dữ liệu/style khác.
    """
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)

    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            cell.alignment = alignment


def _center_all_report_tables(wb, data_rows: int, monthly_rows: int) -> None:
    """
    Căn giữa nội dung tất cả bảng trong workbook.
    Không thêm/xóa dòng, không đổi dữ liệu, không đổi công thức, không đổi chart.
    """

    # 1) Căn giữa các Excel Table có sẵn trong workbook.
    for ws in wb.worksheets:
        for table in ws.tables.values():
            table_ref = table.ref if hasattr(table, "ref") else str(table)
            _center_range(ws, table_ref)

    # 2) Căn giữa các vùng bảng được ghi động bằng code.
    analysis_monthly_end = 17 + max(monthly_rows, 5)
    top_section_row = 26 if monthly_rows <= 5 else 18 + monthly_rows + 3
    top_bottom_end = top_section_row + 6

    manual_table_ranges = {
        "Executive Summary": [
            "A17:L21",
        ],
        "Phân tích": [
            "A4:D13",
            f"A17:I{analysis_monthly_end}",
            f"A{top_section_row}:I{top_bottom_end}",
        ],
        "Dữ liệu sạch": [
            f"A1:U{data_rows + 1}",
        ],
        "Data Issues": [
            "A1:E5",
        ],
        "Change Log": [
            "A1:D9",
        ],
        "Phương pháp": [
            "A1:C7",
        ],
    }

    for sheet_name, ranges in manual_table_ranges.items():
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        for cell_range in ranges:
            _center_range(ws, cell_range)
