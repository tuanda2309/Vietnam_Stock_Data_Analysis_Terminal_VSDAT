import os
from flask import Flask, jsonify, send_file
from flask_cors import CORS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.cell.text import InlineFont
from openpyxl.styles.colors import Color
from io import BytesIO
from datetime import datetime, timedelta
import pandas as pd
import numpy as np

try:
    from vnstock3 import Vnstock
except ImportError:
    from vnstock import Vnstock

app = Flask(__name__)


CORS(app, resources={r"/*": {"origins": "*"}}, supports_credentials=True)

DEVELOPER_EMAIL = os.environ.get('DEVELOPER_EMAIL', 'doananhtuan77qn@gmail.com')
DEVELOPER_PHONE = os.environ.get('DEVELOPER_PHONE', '0399848871')
DEVELOPER_GITHUB = os.environ.get('DEVELOPER_GITHUB', 'https://github.com/tuanda2309')
DEVELOPER_WEB = os.environ.get('DEVELOPER_WEB', 'https://vsdat-frontend.onrender.com')

def calculate_rsi(series, period=14):
    delta = series.diff()
    gain = delta.clip(lower=0)
    loss = (-delta).clip(lower=0)
    avg_gain = gain.ewm(alpha=1/period, adjust=False).mean()
    avg_loss = loss.ewm(alpha=1/period, adjust=False).mean()
    rs = np.where(avg_loss != 0, avg_gain / avg_loss, np.nan)
    rsi = np.where(avg_loss == 0, 100, 100 - (100 / (1 + rs)))
    return np.where((avg_gain == 0) & (avg_loss == 0), 50, rsi)

def calculate_macd(series, fast=12, slow=26, signal=9):
    ema_fast = series.ewm(span=fast, adjust=False).mean()
    ema_slow = series.ewm(span=slow, adjust=False).mean()
    macd_line = ema_fast - ema_slow
    signal_line = macd_line.ewm(span=signal, adjust=False).mean()
    hist = macd_line - signal_line
    return macd_line, signal_line, hist

def calculate_bollinger_bands(series, window=20, num_sd=2):
    ma = series.rolling(window=window).mean()
    sd = series.rolling(window=window).std()
    upper_band = ma + (sd * num_sd)
    lower_band = ma - (sd * num_sd)
    return upper_band, lower_band

# ==========================================
# API LẤY DỮ LIỆU VÀ PHÂN TÍCH CHỨNG KHOÁN
# ==========================================

@app.route('/stock/<symbol>', methods=['GET'])
def get_stock_analysis(symbol):
    try:
        if not symbol or symbol.strip() == "":
            return jsonify({'error': 'Mã cổ phiếu không được để trống'}), 400

        symbol_upper = symbol.strip().upper()
        stock_api = Vnstock().stock(symbol=symbol_upper, source='VCI')

        current_time = datetime.now()
        start_date = (current_time - timedelta(days=365 * 2)).strftime('%Y-%m-%d')
        end_date = current_time.strftime('%Y-%m-%d')

        df = stock_api.quote.history(start=start_date, end=end_date, interval='1D')
        
        if df is None or df.empty:
            return jsonify({'error': f'Không tìm thấy dữ liệu giao dịch cho mã: {symbol_upper}'}), 442

        # 1. Sắp xếp chuỗi thời gian tăng dần để đảm bảo cấu hình chỉ mục (Index) chính xác
        df['time'] = pd.to_datetime(df['time'])
        df = df.sort_values('time').reset_index(drop=True)

        # Ép kiểu dữ liệu số và quy chuẩn về đơn vị VNĐ tuyệt đối
        numeric_cols = ['open', 'high', 'low', 'close', 'volume']
        for col in numeric_cols:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors='coerce')
                if col != 'volume':
                    df[col] = df[col] * 1000

        # 2. Thu thập giá đóng cửa phiên thời gian thực (Realtime Price)
        realtime_price = None
        has_intraday = False
        try:
            intraday = stock_api.quote.intraday()
            if intraday is not None and not intraday.empty:
                realtime_price = float(intraday.iloc[-1]['price']) * 1000
                has_intraday = True
        except Exception:
            pass

        # 3. LOGIC XÁC ĐỊNH GIÁ THAM CHIẾU PHIÊN TRƯỚC AN TOÀN TUYỆT ĐỐI ---
        last_row_date = df.iloc[-1]['time'].date()
        today_date = current_time.date()

        if last_row_date == today_date or has_intraday:
            previous_close = float(df.iloc[-2]['close'])
            if realtime_price is None:
                realtime_price = float(df.iloc[-1]['close'])
        else:
            previous_close = float(df.iloc[-1]['close'])
            if realtime_price is None:
                realtime_price = previous_close

        price_change = realtime_price - previous_close
        percent_change = (price_change / previous_close) * 100

        if price_change > 0:
            status, color = "up", "green"
        elif price_change < 0:
            status, color = "down", "red"
        else:
            status, color = "unchanged", "yellow"

        # --- CONSOLE DEBUG BACKEND ---
        print(f"\n=== DEBUG LOG CHỈ SỐ MÃ {symbol_upper} ===")
        print(f"Giá hiện tại (Current): {realtime_price} VNĐ")
        print(f"Giá phiên trước (Previous Close): {previous_close} VNĐ")
        print(f"Chênh lệch (Change): {price_change} VNĐ | Phần trăm: {round(percent_change, 2)}%")
        print(f"Trạng thái điều hướng màu: {status.upper()} -> {color.upper()}")
        print("=========================================\n")

        # 4. Tính toán hệ thống chỉ báo kỹ thuật nâng cao
        df['MA20'] = df['close'].rolling(window=20).mean()
        df['MA50'] = df['close'].rolling(window=50).mean()
        df['MA100'] = df['close'].rolling(window=100).mean()
        df['MA200'] = df['close'].rolling(window=200).mean()
        df['RSI'] = calculate_rsi(df['close'], period=14)
        
        macd_l, signal_l, _ = calculate_macd(df['close'])
        df['MACD'] = macd_l
        df['Signal_Line'] = signal_l
        
        b_upper, b_lower = calculate_bollinger_bands(df['close'])
        df['Bollinger_Upper'] = b_upper
        df['Bollinger_Lower'] = b_lower

        df['MA20'] = df['MA20'].fillna(df['close'])
        df['RSI'] = df['RSI'].fillna(50)

        latest = df.iloc[-1]
        rsi_val = latest['RSI'] if pd.notna(latest['RSI']) else 50
        ma20_val = latest['MA20'] if pd.notna(latest['MA20']) else realtime_price

        df = df.replace({np.nan: None})
        df['time'] = df['time'].dt.strftime('%Y-%m-%d')
        chart_data = df.tail(120).to_dict(orient='records')

        return jsonify({
            'symbol': symbol_upper,
            'currentPrice': realtime_price,
            'referencePrice': previous_close,
            'previousClose': previous_close,
            'priceChange': round(price_change, 2),
            'percentChange': round(percent_change, 2),
            'status': status,
            'color': color,
            'rsi': round(float(rsi_val), 2),
            'ma20': round(float(ma20_val), 2),
            'data': chart_data
        }), 200

    except Exception as e:
        return jsonify({'error': f'Lỗi hệ thống Máy chủ Backend: {str(e)}'}), 500

# ==========================================
# API XUẤT FILE BÁO CÁO EXCEL
# ==========================================

@app.route('/export/<symbol>/<start_date>/<end_date>', methods=['GET'])
def export_excel(symbol, start_date, end_date):
    try:
        if not symbol or symbol.strip() == "":
            return jsonify({'error': 'Mã cổ phiếu không được để trống'}), 400

        symbol_upper = symbol.strip().upper()
        
        try:
            start_obj = datetime.strptime(start_date, '%Y-%m-%d')
            end_obj = datetime.strptime(end_date, '%Y-%m-%d')
        except ValueError:
            return jsonify({'error': 'Định dạng ngày không hợp lệ. Vui lòng dùng YYYY-MM-DD'}), 400

        if start_obj > end_obj:
            return jsonify({'error': 'Ngày bắt đầu không được lớn hơn ngày kết thúc'}), 400

        if start_obj > datetime.now() or end_obj > datetime.now():
            return jsonify({'error': 'Hệ thống không hỗ trợ trích xuất dữ liệu ngày tương lai'}), 400

        stock_api = Vnstock().stock(symbol=symbol_upper, source='VCI')
        df = stock_api.quote.history(start=start_date, end=end_date, interval='1D')

        if df is None or df.empty:
            return jsonify({'error': f'Không có dữ liệu lịch sử để xuất file cho mã {symbol_upper}'}), 442

        df['time'] = pd.to_datetime(df['time'])
        df = df.sort_values('time').reset_index(drop=True)

        numeric_cols = ['open', 'high', 'low', 'close', 'volume']
        for col in numeric_cols:
            df[col] = pd.to_numeric(df[col], errors='coerce')
            if col != 'volume':
                df[col] = df[col] * 1000

        df['Change'] = df['close'].diff().fillna(0)
        df['PercentChange'] = (df['close'].pct_change() * 100).fillna(0)

        df = df[df['time'] >= pd.to_datetime(start_date)].reset_index(drop=True)

        wb = Workbook()
        ws = wb.active
        ws.title = symbol_upper

        normal_font = Font(name='Arial', size=8)
        header_font = Font(name='Arial', size=8, color='FFFFFF', bold=True)
        blue_fill = PatternFill(start_color='0066CC', end_color='0066CC', fill_type='solid')
        center = Alignment(horizontal='center', vertical='center')

        inline_bold = InlineFont(rFont='Arial', sz=8, b=True)
        inline_normal = InlineFont(rFont='Arial', sz=8, b=False)
        inline_link = InlineFont(rFont='Arial', sz=8, b=False, u='single', color=Color(rgb='0000FF'))

        # Thiết lập khối thông tin nhà phát triển (Dòng 1 đến Dòng 4)
        ws['A1'] = CellRichText([TextBlock(inline_bold, 'Email: '), TextBlock(inline_normal, DEVELOPER_EMAIL)])
        ws['A2'] = CellRichText([TextBlock(inline_bold, 'Phone: '), TextBlock(inline_normal, DEVELOPER_PHONE)])
        
        ws['A3'] = CellRichText([TextBlock(inline_bold, 'GitHub: '), TextBlock(inline_link, 'Tại đây')])
        ws['A3'].hyperlink = DEVELOPER_GITHUB
        ws['A3'].style = 'Hyperlink'

        ws['A4'] = CellRichText([TextBlock(inline_bold, 'Web: '), TextBlock(inline_link, 'Tại đây')])  
        ws['A4'].hyperlink = DEVELOPER_WEB                                                          
        ws['A4'].style = 'Hyperlink'                                                                

        # Cấu hình vị trí hàng tiêu đề chính (Dời từ dòng 4 xuống dòng 5)
        headers = ['NGÀY', 'GIÁ MỞ CỬA', 'GIÁ CAO NHẤT', 'GIÁ THẤP NHẤT', 'GIÁ ĐÓNG CỬA', 'THAY ĐỔI GIÁ', '% THAY ĐỔI', 'KHỐI LƯỢNG']
        header_row = 5

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.value = header
            cell.fill = blue_fill
            cell.font = header_font
            cell.alignment = center

        # Ghi mảng dữ liệu lịch sử (Dời chỉ mục bắt đầu từ dòng 5 xuống dòng 6)
        data_start_row = 6
        for row_num, (_, row) in enumerate(df.iterrows(), data_start_row):
            values = [
                row['time'].strftime('%d/%m/%Y'),
                float(row['open']), float(row['high']), float(row['low']), float(row['close']),
                float(row['Change']), float(row['PercentChange']), float(row['volume'])
            ]

            for col_num, value in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.font = normal_font
                cell.alignment = center

                if col_num in [2, 3, 4, 5, 6]:
                    cell.number_format = '#,##0.00'
                elif col_num == 7:
                    cell.number_format = '0.00'
                elif col_num == 8:
                    cell.number_format = '#,##0'

        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value and not isinstance(cell.value, CellRichText):
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = max_length + 5
            ws.column_dimensions[column_letter].width = adjusted_width
            
        ws.column_dimensions['A'].width = 14

        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        # BỔ SUNG EXPOSE HEADERS ĐỂ FRONTEND CÓ THỂ ĐỌC ĐƯỢC THÔNG TIN FILE (Tránh lỗi CORS khi tải file)
        response = send_file(
            excel_file,
            download_name=f'{symbol_upper}_{start_date}_{end_date}.xlsx',
            as_attachment=True,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
        return response

    except Exception as e:
        return jsonify({'error': f'Lỗi xuất tệp Excel hệ thống: {str(e)}'}), 500

# ĐỂ CHẠY ĐƯỢC TRÊN RENDER KHI DÙNG GUNICORN KHÔNG CẦN CHẠY BẰNG LỆNH IF __MAIN__ NÀY.
# TUY NHIÊN VẪN GIỮ ĐỂ BẠN CHẠY LOCALHOST NẾU MUỐN.
if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
